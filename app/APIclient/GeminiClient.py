import base64
import html
import logging
import os
import zipfile
from io import BytesIO
from aiohttp_socks import ProxyConnector
import aiohttp

try:
    from dotenv import load_dotenv
    load_dotenv()
except Exception:
    pass

try:
    from ..config import GEMINI_API_KEY as _CONFIG_GEMINI_KEY
except Exception:
    _CONFIG_GEMINI_KEY = None

from ..prompts import DEFAULT_TEXT_MODEL, TEXT_CHAT_MODELS

GEMINI_API_BASE_URL = "https://generativelanguage.googleapis.com/v1beta"

# Предпочтительный порядок моделей Gemini (сначала быстрые и доступные по квотам)
PREFERRED_GEMINI_MODELS = (
    "gemini-2.5-flash",
    "gemini-flash-latest",
    "gemini-2.5-flash-lite",
    "gemini-3-flash-preview",
    "gemini-2.5-pro",
    "gemini-pro-latest",
    "gemini-1.5-flash",
)

# Ключевые слова моделей, которые НЕ предназначены для обычного текстового чата
_EXCLUDED_MODEL_SUBSTRINGS = (
    "tts",
    "image",
    "imagen",
    "veo",
    "embedding",
    "aqa",
    "robotics",
    "computer-use",
    "audio",
    "learnlm",
    "gemma",
)

MAX_DIALOG_MESSAGES = 24
_USER_SESSIONS: dict[int, dict] = {}
_CACHED_AVAILABLE_MODELS: list[str] | None = None


def get_gemini_api_key() -> str:
    key = (
        _CONFIG_GEMINI_KEY
        or os.getenv("GEMINI_API_KEY")
        or os.getenv("GOOGLE_API_KEY")
        or ""
    ).strip()
    if not key:
        raise RuntimeError(
            "Не задан GEMINI_API_KEY. Добавьте GEMINI_API_KEY=ваш_ключ в файл .env или в config.py "
            "(получить ключ можно в Google AI Studio: https://aistudio.google.com/apikey)."
        )
    return key


def get_gemini_proxy() -> str | None:
    proxy = (os.getenv("GEMINI_PROXY") or os.getenv("HTTPS_PROXY") or "").strip()
    return proxy or None


def get_user_session(user_id: int) -> dict:
    if user_id not in _USER_SESSIONS:
        _USER_SESSIONS[user_id] = {
            "model": DEFAULT_TEXT_MODEL,
            "remember_dialog": True,
            "active_dialog_id": 1,
            "next_dialog_id": 2,
            "dialogs": {
                1: {
                    "title": "Диалог #1",
                    "messages": [],
                }
            },
        }
    return _USER_SESSIONS[user_id]


def get_user_text_model(user_id: int) -> str:
    session = get_user_session(user_id)
    model = session.get("model", DEFAULT_TEXT_MODEL)
    return model if model in TEXT_CHAT_MODELS else DEFAULT_TEXT_MODEL


def set_user_text_model(user_id: int, model_key: str) -> str:
    session = get_user_session(user_id)
    if model_key in TEXT_CHAT_MODELS:
        session["model"] = model_key
    return session["model"]


def is_remember_dialog_enabled(user_id: int) -> bool:
    return bool(get_user_session(user_id).get("remember_dialog", True))


def toggle_remember_dialog(user_id: int) -> bool:
    session = get_user_session(user_id)
    session["remember_dialog"] = not bool(session.get("remember_dialog", True))
    return session["remember_dialog"]


def get_active_dialog(user_id: int) -> dict:
    session = get_user_session(user_id)
    active_id = session["active_dialog_id"]
    if active_id not in session["dialogs"]:
        session["dialogs"][active_id] = {"title": f"Диалог #{active_id}", "messages": []}
    return session["dialogs"][active_id]


def create_new_user_dialog(user_id: int) -> int:
    session = get_user_session(user_id)
    new_id = session["next_dialog_id"]
    session["next_dialog_id"] += 1
    session["dialogs"][new_id] = {
        "title": f"Диалог #{new_id}",
        "messages": [],
    }
    session["active_dialog_id"] = new_id
    return new_id


def switch_user_dialog(user_id: int, dialog_id: int) -> bool:
    session = get_user_session(user_id)
    if dialog_id in session["dialogs"]:
        session["active_dialog_id"] = dialog_id
        return True
    return False


def clear_active_dialog_messages(user_id: int) -> None:
    dialog = get_active_dialog(user_id)
    dialog["messages"] = []


async def _discover_available_models(
    session: aiohttp.ClientSession,
    api_key: str,
    proxy: str | None = None,
) -> list[str]:
    """
    Запрашивает у Google AI Studio список реально доступных для данного ключа моделей
    через GET /v1beta/models и сортирует их по приоритету.
    """
    global _CACHED_AVAILABLE_MODELS
    if _CACHED_AVAILABLE_MODELS:
        return _CACHED_AVAILABLE_MODELS

    url = f"{GEMINI_API_BASE_URL}/models?key={api_key}"
    try:
        async with session.get(url) as response:
            if response.status == 200:
                data = await response.json(content_type=None)
                models_raw = data.get("models") or []
                discovered: list[str] = []
                for m in models_raw:
                    methods = m.get("supportedGenerationMethods") or []
                    if "generateContent" not in methods:
                        continue
                    full_name = m.get("name", "")
                    short_name = full_name.replace("models/", "").strip()
                    if not short_name.startswith("gemini"):
                        continue
                    if any(excl in short_name.lower() for excl in _EXCLUDED_MODEL_SUBSTRINGS):
                        continue
                    discovered.append(short_name)

                if discovered:
                    ordered: list[str] = []
                    for pref in PREFERRED_GEMINI_MODELS:
                        if pref in discovered and pref not in ordered:
                            ordered.append(pref)
                    for name in discovered:
                        if name not in ordered:
                            ordered.append(name)
                    _CACHED_AVAILABLE_MODELS = ordered
                    logging.info("Доступные модели Gemini для ключа: %s", ordered[:6])
                    return ordered
    except Exception as exc:
        logging.debug("Не удалось получить список моделей через GET /v1beta/models: %s", exc)

    return list(PREFERRED_GEMINI_MODELS)


async def _call_gemini_generate_content(
    contents: list[dict],
    system_instruction: str | None = None,
    temperature: float = 0.7,
    max_output_tokens: int = 8192,
    timeout: float = 180.0,
) -> str:
    """Выполняет запрос к Gemini API через отдельный SOCKS5-прокси."""

    api_key = get_gemini_api_key()
    proxy = get_gemini_proxy()

    payload: dict = {
        "contents": contents,
        "generationConfig": {
            "temperature": temperature,
            "maxOutputTokens": max_output_tokens,
        },
    }

    if system_instruction:
        payload["systemInstruction"] = {
            "parts": [{"text": system_instruction}],
        }

    client_timeout = aiohttp.ClientTimeout(
        total=timeout,
        connect=15,
        sock_read=timeout,
    )
    errors: list[str] = []

    # Прокси используется только этой HTTP-сессией Gemini.
    connector = ProxyConnector.from_url(proxy) if proxy else None

    async with aiohttp.ClientSession(
        timeout=client_timeout,
        connector=connector,
    ) as session:
        # Получаем доступные модели через ту же сессию и прокси.
        models_to_try = await _discover_available_models(
            session,
            api_key,
        )

        for gemini_model in models_to_try[:5]:
            url = (
                f"{GEMINI_API_BASE_URL}/models/"
                f"{gemini_model}:generateContent?key={api_key}"
            )

            try:
                async with session.post(
                    url,
                    json=payload,
                ) as response:
                    data = await response.json(content_type=None)

                    if response.status >= 400:
                        err_msg = (
                            (data.get("error") or {}).get("message")
                            if isinstance(data, dict)
                            else str(data)
                        ) or f"HTTP {response.status}"

                        logging.warning(
                            "Модель %s вернула HTTP %s: %s",
                            gemini_model,
                            response.status,
                            err_msg,
                        )

                        errors.append(
                            f"{gemini_model} (HTTP {response.status}): "
                            f"{err_msg}"
                        )

                        lower_msg = err_msg.lower()

                        if (
                            "api key not valid" in lower_msg
                            or "api_key_invalid" in lower_msg
                            or "user location is not supported" in lower_msg
                        ):
                            break

                        continue

                    candidates = data.get("candidates") or []

                    if not candidates:
                        block_reason = (
                            (data.get("promptFeedback") or {}).get(
                                "blockReason"
                            )
                        )
                        errors.append(
                            f"{gemini_model}: пустой список candidates "
                            f"(blockReason={block_reason})"
                        )
                        continue

                    parts = (
                        (candidates[0].get("content") or {}).get("parts")
                        or []
                    )
                    text_chunks = [
                        part.get("text", "")
                        for part in parts
                        if part.get("text")
                    ]
                    answer = "".join(text_chunks).strip()

                    if answer:
                        return answer

                    errors.append(
                        f"{gemini_model}: получен пустой текст ответа"
                    )

            except Exception as exc:
                logging.warning(
                    "Ошибка соединения при вызове %s: %s",
                    gemini_model,
                    exc,
                )
                errors.append(f"{gemini_model}: {exc}")

    first_error = errors[0] if errors else "Неизвестная ошибка"

    if "user location is not supported" in first_error.lower():
        raise RuntimeError(
            "Gemini отклонил запрос из-за региона подключения. "
            "Проверьте страну и работоспособность GEMINI_PROXY."
        )

    raise RuntimeError(
        f"Не удалось получить ответ от Gemini API: {first_error}"
    )


async def generate_gemini_chat_response(
    user_id: int,
    user_text: str,
    image_bytes: bytes | None = None,
    mime_type: str = "image/jpeg",
) -> str:
    """
    Отвечает пользователю в обычном ИИ-чате от лица выбранной модели (GPT 5 Luna / Sol / Terra / Claude / Grok),
    используя Gemini API и сохраняя историю диалога (если включено запоминание).
    """
    model_key = get_user_text_model(user_id)
    model_cfg = TEXT_CHAT_MODELS[model_key]
    system_prompt = model_cfg["system_prompt"]

    remember = is_remember_dialog_enabled(user_id)
    active_dialog = get_active_dialog(user_id)
    history: list[dict] = active_dialog["messages"] if remember else []

    contents: list[dict] = []
    for item in history[-MAX_DIALOG_MESSAGES:]:
        contents.append({
            "role": "user" if item["role"] == "user" else "model",
            "parts": [{"text": item["text"]}],
        })

    current_parts: list[dict] = []
    if image_bytes:
        current_parts.append({
            "inline_data": {
                "mime_type": mime_type,
                "data": base64.b64encode(image_bytes).decode("ascii"),
            }
        })
    current_parts.append({"text": user_text or "Опиши это изображение и ответь по его содержимому."})
    contents.append({"role": "user", "parts": current_parts})

    answer = await _call_gemini_generate_content(
        contents=contents,
        system_instruction=system_prompt,
        temperature=0.7,
        max_output_tokens=8192,
    )

    if remember:
        active_dialog["messages"].append({"role": "user", "text": user_text or "[Изображение]"})
        active_dialog["messages"].append({"role": "model", "text": answer})
        if len(active_dialog["messages"]) > MAX_DIALOG_MESSAGES:
            active_dialog["messages"] = active_dialog["messages"][-MAX_DIALOG_MESSAGES:]
        if len(active_dialog["messages"]) == 2 and user_text:
            short_title = user_text.strip().replace("\n", " ")[:30]
            active_dialog["title"] = short_title + ("…" if len(user_text.strip()) > 30 else "")

    return answer


async def generate_gemini_service_response(
    system_prompt: str,
    user_text: str,
    image_bytes: bytes | None = None,
    mime_type: str = "image/jpeg",
) -> str:
    """Выполняет специализированный запрос к Gemini для учебных/рабочих сервисов и распознавания фото."""
    parts: list[dict] = []
    if image_bytes:
        parts.append({
            "inline_data": {
                "mime_type": mime_type,
                "data": base64.b64encode(image_bytes).decode("ascii"),
            }
        })
    parts.append({"text": user_text or "Выполни задание по изображению."})

    return await _call_gemini_generate_content(
        contents=[{"role": "user", "parts": parts}],
        system_instruction=system_prompt,
        temperature=0.6,
        max_output_tokens=8192,
    )


def _build_minimal_docx_fallback(topic: str, work_type_title: str, text_content: str) -> bytes:
    """Резервная сборка .docx без внешних библиотек (на случай если python-docx не установлен)."""
    paragraphs = [work_type_title.upper(), f"Тема: «{topic.strip()}»", ""]
    paragraphs.extend(line.replace("**", "") for line in text_content.splitlines())

    body_xml = "".join(
        f"<w:p><w:r><w:t>{html.escape(p)}</w:t></w:r></w:p>"
        for p in paragraphs
    )
    document_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f"<w:body>{body_xml}</w:body></w:document>"
    )
    content_types_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
        "</Types>"
    )
    rels_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
        'Target="word/document.xml"/>'
        "</Relationships>"
    )
    buf = BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", content_types_xml)
        zf.writestr("_rels/.rels", rels_xml)
        zf.writestr("word/document.xml", document_xml)
    return buf.getvalue()


def build_academic_docx_bytes(
    topic: str,
    work_type_title: str,
    text_content: str,
) -> bytes:
    """
    Собирает аккуратно оформленный Word-файл (.docx) с текстом сгенерированной работы
    (Эссе, Курсовая работа, Дипломная работа, Реферат).
    """
    try:
        import docx
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt, RGBColor
    except ImportError:
        return _build_minimal_docx_fallback(topic, work_type_title, text_content)

    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    header_p = doc.add_paragraph()
    header_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_run = header_p.add_run(work_type_title.upper())
    header_run.font.name = "Times New Roman"
    header_run.font.size = Pt(16)
    header_run.font.bold = True
    header_p.paragraph_format.space_after = Pt(6)

    topic_p = doc.add_paragraph()
    topic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    topic_run = topic_p.add_run(f"Тема: «{topic.strip()}»")
    topic_run.font.name = "Times New Roman"
    topic_run.font.size = Pt(14)
    topic_run.font.bold = True
    topic_p.paragraph_format.space_after = Pt(18)

    for raw_line in text_content.splitlines():
        line = raw_line.strip()
        if not line:
            continue

        clean_line = line.replace("**", "")

        if clean_line.startswith("### "):
            p = doc.add_paragraph()
            run = p.add_run(clean_line[4:].strip())
            run.font.name = "Times New Roman"
            run.font.size = Pt(13)
            run.font.bold = True
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
        elif clean_line.startswith("## "):
            p = doc.add_paragraph()
            run = p.add_run(clean_line[3:].strip())
            run.font.name = "Times New Roman"
            run.font.size = Pt(14)
            run.font.bold = True
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
        elif clean_line.startswith("# "):
            p = doc.add_paragraph()
            run = p.add_run(clean_line[2:].strip())
            run.font.name = "Times New Roman"
            run.font.size = Pt(15)
            run.font.bold = True
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
        else:
            p = doc.add_paragraph()
            run = p.add_run(clean_line)
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.first_line_indent = Inches(0.5)
            p.paragraph_format.space_after = Pt(6)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


async def check_gemini_api() -> tuple[bool, str]:
    """Проверяет подключение к Gemini API с текущими настройками прокси."""
    try:
        answer = await _call_gemini_generate_content(
            contents=[
                {
                    "role": "user",
                    "parts": [
                        {"text": "Ответь одним словом: OK"}
                    ],
                }
            ],
            temperature=0,
            max_output_tokens=20,
            timeout=30,
        )

        return True, answer

    except Exception as exc:
        logging.exception("Проверка подключения Gemini API завершилась ошибкой")
        return False, str(exc)    